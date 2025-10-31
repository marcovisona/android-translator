import codecs
import collections
import csv
import os
import sys
from xml.dom import minidom
import openpyxl

def escapeAndroidChar(text):
    text = text.replace('\'', "\\'")
    return text


def get_original_key_order(xml_file_path):
    """
    Read the existing XML file and extract the order of keys (both regular strings and arrays).
    Returns an ordered list of keys as they appear in the original XML.
    """
    if not os.path.exists(xml_file_path):
        return []
    
    key_order = []
    try:
        xmldoc = minidom.parse(xml_file_path)
        rootNode = xmldoc.getElementsByTagName("resources")
        if len(rootNode) == 1:
            nodeList = rootNode[0].childNodes
            for n in nodeList:
                if hasattr(n, 'attributes') and n.attributes is not None:
                    tag = n.tagName
                    if tag == 'string':
                        key = n.attributes['name'].nodeValue
                        key_order.append(key)
                    elif tag == 'string-array':
                        name = n.attributes['name'].nodeValue
                        itemList = n.getElementsByTagName("item")
                        for idx in range(len(itemList)):
                            key = str(name) + "," + str(idx)
                            key_order.append(key)
    except Exception as e:
        print(f"Warning: Could not read original XML file {xml_file_path}: {e}")
    
    return key_order


# pathToString = raw_input("Path to CSV file:")
# outputFolder = raw_input("Path to Android project (output):")
# defaultLanguage = raw_input("Default langage ISO-639-1 code. (write en if your default langage is english):")

pathToExportString = sys.argv[1]
# pathToString = sys.argv[2]
outputFolder = sys.argv[2]
defaultLanguage = "en"

outputFolder = os.path.join(outputFolder, "res")


def read_xlsx(filename):
    content = dict()
    wb = openpyxl.load_workbook(filename)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))

    if len(rows) <= 0:
        return content

    header_tmp = rows[0]
    langList = []
    for lang in header_tmp[1:]:
        if lang:
            content[lang] = dict()
            langList.append(lang)
    for l in rows[1:]:
        key = l[0]
        for idx, item in enumerate(l[1:]):
            if item:
                content[langList[idx]][key] = escapeAndroidChar(item)
            else:
                print("Empty item for " + langList[idx])
    return content

def read_csv(filename):
    # Open the CSV file in read mode
    rows = []
    with open(filename, 'r') as csvfile:
        # Create a CSV reader object
        reader = csv.reader(csvfile)

        # Read and process each row
        for row in reader:
            rows.append(row)

    return rows


def read_file(file_name):
    content = dict()
    lines = read_csv(file_name)
    if len(lines) <= 0:
        return content

    header_tmp = lines[0]
    langList = []
    for lang in header_tmp[1:]:
        if lang:
            content[lang] = dict()
            langList.append(lang)
    for l in lines[1:]:
        key = l[0]
        for idx, item in enumerate(l[1:]):
            if item:
                # 3) Escape character
                content[langList[idx]][key] = escapeAndroidChar(item)
            else:
                print("Empty item for " + langList[idx])
    return content

languageExportDict = read_xlsx(pathToExportString)
# languageExportDict = read_file(pathToExportString)
xmldict = dict()


def load_strings_arr(lang):
    with codecs.open(lang, 'r', "utf-8") as f:
        return [x.strip("\n") for x in f.readlines()]


def add(collection, is_ordered=False):
    """
    Add string elements to the XML document.
    
    Args:
        collection: Either a dict or a list of keys to process
        is_ordered: If True, collection is a list of keys in specific order.
                   If False, collection is a dict to iterate over.
    """
    if is_ordered:
        # collection is a list of keys in the desired order
        keys_to_process = collection
    else:
        # collection is a dict, iterate over its keys
        keys_to_process = collection
    
    current_array_node = None
    current_array_name = None
    
    for key in keys_to_process:
        if key is None:
            continue

        # Check if this key exists in stringsDict
        if not stringsDict.get(key):
            continue

        if ',' in key:
            # This is a string-array item
            array_name = key[:key.rfind(',')]
            
            # If this is a new array or the first item (,0), create the array node
            if current_array_name != array_name:
                current_array_node = doc.createElement("string-array")
                current_array_node.setAttribute("name", array_name)
                rootNode.appendChild(current_array_node)
                current_array_name = array_name
            
            # Add item to the current array
            node = doc.createElement("item")
            node.appendChild(doc.createTextNode(stringsDict[key]))
            current_array_node.appendChild(node)
        else:
            # Regular string element
            current_array_name = None  # Reset array tracking
            node = doc.createElement("string")
            node.setAttribute("name", key)
            node.appendChild(doc.createTextNode(stringsDict[key]))
            rootNode.appendChild(node)


for lang in languageExportDict:
    stringsDict = collections.OrderedDict(languageExportDict[lang])
    stringsExportDict = languageExportDict[lang]
    doc = minidom.Document()
    xmldict[lang] = doc
    rootNode = doc.createElement("resources")
    doc.appendChild(rootNode)

    # Determine the output path for this language to read original ordering
    folderName = "values" if lang == defaultLanguage else "values-" + lang
    langFolder = os.path.join(outputFolder, folderName)
    stringPath = os.path.join(langFolder, "strings.xml")
    
    # Get the original key order from the existing XML file
    original_key_order = get_original_key_order(stringPath)
    
    # Separate keys into: existing (in original order) and new (not in original XML)
    existing_keys = []
    new_keys = []
    
    for key in stringsDict.keys():
        if key in original_key_order:
            existing_keys.append(key)
        else:
            new_keys.append(key)
    
    # Sort existing keys by their original position
    existing_keys.sort(key=lambda k: original_key_order.index(k))
    
    # Add existing keys first (preserving original order)
    add(existing_keys, is_ordered=True)
    
    # Add new keys at the end (in the order they appear in CSV)
    add(new_keys, is_ordered=True)

# 3) Write XML for each langage in the correct directory structure
for lang in languageExportDict:
    folderName = "values"
    if lang != defaultLanguage:
        folderName = "values-" + lang
    langFolder = os.path.join(outputFolder, folderName)
    if not os.path.exists(langFolder):
        os.makedirs(langFolder)
    stringPath = os.path.join(langFolder, "strings.xml")
    with open(stringPath, 'wb') as f:
        xmlContain = xmldict[lang].toprettyxml(encoding="utf-8", indent='    ')
        # print(xmlContain)
        f.write(xmlContain)
